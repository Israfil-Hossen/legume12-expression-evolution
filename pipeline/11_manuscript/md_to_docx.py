#!/usr/bin/env python3
"""Convert the rendered manuscript (Markdown subset) to a submission-style Word file:
Times New Roman 12 pt, double spacing, continuous line numbers, page numbers not required.
Supports headings (#, ##), paragraphs, **bold**, *italic*, ^superscript^, subscripts written as
symbol_word (rho_between -> rho with subscript 'between'), bullet lists and pipe tables.
Usage: python pipeline/11_manuscript/md_to_docx.py manuscript_mpe/manuscript.md manuscript_mpe/manuscript.docx
"""
import re, sys
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

src, dst = sys.argv[1], sys.argv[2]
lines = open(src, encoding="utf-8").read().split("\n")
doc = Document()
st = doc.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(12)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
st.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
st.paragraph_format.space_after = Pt(0)
for s in ("Heading 1", "Heading 2"):
    h = doc.styles[s]; h.font.name = "Times New Roman"; h.font.color.rgb = None
    h.font.size = Pt(14 if s == "Heading 1" else 12); h.font.bold = True; h.font.italic = s == "Heading 2"
sec = doc.sections[0]
for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
    setattr(sec, side, Cm(2.5))
ln = OxmlElement("w:lnNumType"); ln.set(qn("w:countBy"), "1"); ln.set(qn("w:restart"), "continuous")
sec._sectPr.append(ln)

SUB = re.compile(r"(ρ|σ²|d|x)_([A-Za-z]+)")
TOK = re.compile(r"(\*\*(?:[^*]|\*[^*]+\*)+\*\*|\*[^*]+\*|\^[^^]+\^|`[^`]+`)")


def add_runs(par, text, bold=False):
    for part in TOK.split(text):
        if not part:
            continue
        b, i, sup, txt = bold, False, False, part
        if part.startswith("**") and part.endswith("**"):
            add_runs(par, part[2:-2], bold=True)          # bold may contain italic species names
            continue
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            i, txt = True, part[1:-1]
        elif part.startswith("^") and part.endswith("^"):
            sup, txt = True, part[1:-1]
        elif part.startswith("`") and part.endswith("`"):
            txt = part[1:-1]
        txt = txt.replace("\\*", "*")
        pos = 0
        for m in SUB.finditer(txt):
            r = par.add_run(txt[pos:m.start()] + m.group(1)); r.bold, r.italic, r.font.superscript = b, i, sup
            r = par.add_run(m.group(2)); r.bold, r.italic = b, i; r.font.subscript = True
            pos = m.end()
        r = par.add_run(txt[pos:]); r.bold, r.italic, r.font.superscript = b, i, sup


i = 0
while i < len(lines):
    line = lines[i]
    if not line.strip():
        i += 1; continue
    if line.startswith("|"):
        rows = []
        while i < len(lines) and lines[i].startswith("|"):
            cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
            if not all(re.fullmatch(r":?-+:?", c) for c in cells):
                rows.append(cells)
            i += 1
        t = doc.add_table(rows=len(rows), cols=len(rows[0])); t.style = "Table Grid"
        for r_, row in enumerate(rows):
            for c_, cell in enumerate(row):
                p = t.cell(r_, c_).paragraphs[0]; p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
                add_runs(p, cell, bold=(r_ == 0))
                for run in p.runs:
                    run.font.size = Pt(9)
        doc.add_paragraph()
        continue
    if line.startswith("!["):                                   # embedded figure: ![](path), full text width
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        doc.add_picture(re.search(r"\]\(([^)]+)\)", line).group(1), width=Cm(16))
        pic = doc.paragraphs[-1]; pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pic.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE; pic.paragraph_format.keep_with_next = True
        i += 1; continue
    if line.startswith("# "):
        add_runs(doc.add_heading("", level=1), line[2:].strip())
    elif line.startswith("## "):
        add_runs(doc.add_heading("", level=2), line[3:].strip())
    elif line.startswith("- "):
        add_runs(doc.add_paragraph(style="List Bullet"), line[2:].strip())
    else:
        add_runs(doc.add_paragraph(), line.strip())
    i += 1
import datetime, re as _re
cp = doc.core_properties
first = open(src, encoding="utf-8").readline().strip().strip("*")
now = datetime.datetime.now(datetime.timezone.utc).replace(microsecond=0, tzinfo=None)
cp.title = first; cp.author = "Md. Isrfil Hossen"; cp.last_modified_by = "Md. Isrfil Hossen"
cp.created = now; cp.modified = now; cp.revision = 1
cp.comments = ""; cp.keywords = ""; cp.subject = ""; cp.category = ""
doc.save(dst)
print("wrote", dst)
